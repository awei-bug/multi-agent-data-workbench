from __future__ import annotations

import base64
import json
import logging
import time
from datetime import datetime, timezone
from html import escape
from pathlib import Path
from uuid import uuid4

from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel

from backend_py.agnes_client import AgnesChatClient, create_agnes_client
from backend_py.artifact_store import ArtifactStore, create_artifact_store
from backend_py.auth import (
    AuthError,
    AuthUser,
    PERMISSION_ADMIN,
    PERMISSION_READ,
    PERMISSION_WRITE,
    ROLE_ADMIN,
    ROLE_ANALYST,
    ROLE_LOCAL,
    ROLE_VIEWER,
    create_access_token,
    decode_access_token,
    has_permission,
    hash_password,
    is_write_method,
    normalize_role,
    permissions_for_role,
    public_auth_paths,
    user_from_token_payload,
    validate_password,
    validate_username,
    verify_password,
)
from backend_py.connectors import (
    ConnectorError,
    build_connector_data_source,
    decrypt_credentials,
    encrypt_credentials,
    fetch_connector_payload,
    list_connector_types,
    public_connection_view,
    test_connection,
)
from backend_py.data_source_ingest import ingest_data_source
from backend_py.error_tracking import get_error_tracker, init_error_tracking
from backend_py.observability import get_metrics
from backend_py.rate_limit import create_rate_limiter, rate_limit_key
from backend_py.repository import TaskRepository, create_repository, get_workspace_scope, reset_workspace_scope, set_workspace_scope
from backend_py.security import security_headers, validate_production_settings
from backend_py.upload_security import UploadSecurityError, validate_upload
from backend_py.workflow_runner import create_workflow_queue, attach_execution_meta

logger = logging.getLogger("workspace_api")


class IntentRequest(BaseModel):
    request: str = ""
    dataSource: dict | None = None


class ActiveTaskUpdate(BaseModel):
    name: str | None = None
    mode: str | None = None
    status: str | None = None
    summary: str | None = None
    request: str | None = None
    intent: dict | None = None
    snapshot: dict | None = None
    exports: list[dict] | None = None
    risks: list[str] | None = None
    dataSource: dict | None = None


class WorkflowActionRequest(BaseModel):
    action: str


class TemplateDraftRequest(BaseModel):
    name: str | None = None
    request: str | None = None
    summary: str | None = None


class DataSourceIngestRequest(BaseModel):
    fileName: str
    contentBase64: str
    mimeType: str | None = None


class ExportRequest(BaseModel):
    format: str
    report: str | None = None


class TaskCorrectionRequest(BaseModel):
    summary: str


class AuthRegisterRequest(BaseModel):
    username: str
    password: str
    tenantId: str | None = None
    displayName: str | None = None


class AuthLoginRequest(BaseModel):
    username: str
    password: str


class RoleUpdateRequest(BaseModel):
    role: str


class ConnectionUpsertRequest(BaseModel):
    name: str
    connectorType: str
    config: dict = {}
    credentials: dict | None = None


class ConnectionTestRequest(BaseModel):
    connectorType: str | None = None
    config: dict | None = None
    credentials: dict | None = None


class ConnectionImportRequest(BaseModel):
    datasetName: str | None = None


def _audit(repo: TaskRepository, action: str, target: str, payload: dict | None = None) -> None:
    repo.add_audit_event(action, target, payload)


def _add_security_headers(response, *, enable_hsts: bool = False, csp: str | None = None) -> None:
    for key, value in security_headers(enable_hsts=enable_hsts, csp=csp).items():
        response.headers[key] = value


def _error_response(
    request_id: str,
    status_code: int,
    code: str,
    message: str,
    details: dict | None = None,
    extra_headers: dict[str, str] | None = None,
    *,
    enable_hsts: bool = False,
    csp: str | None = None,
):
    response = JSONResponse(
        status_code=status_code,
        content={
            "error": {
                "code": code,
                "message": message,
                "details": details or {},
            },
            "requestId": request_id,
        },
    )
    response.headers["x-request-id"] = request_id
    if extra_headers:
        for key, value in extra_headers.items():
            response.headers[key] = value
    _add_security_headers(response, enable_hsts=enable_hsts, csp=csp)
    return response


def _data_source_context(data_source: dict | None) -> str:
    if not data_source:
        return ""
    preview = data_source.get("textPreview") or json.dumps(data_source.get("previewRows", []), ensure_ascii=False)
    if not preview:
        return ""
    return f"\nUploaded file: {data_source.get('fileName', 'unknown')}\nContent preview:\n{preview[:2000]}"


def _field_names(data_source: dict | None, preferred_type: str | None = None) -> list[str]:
    columns = (data_source or {}).get("columns") or []
    names = []
    for column in columns:
        if preferred_type and column.get("type") != preferred_type:
            continue
        name = column.get("name")
        if name:
            names.append(name)
    return names


def _build_agentic_intent(request_text: str, data_source: dict | None, model_payload: dict | None = None) -> dict:
    data_source = data_source or {}
    is_document = data_source.get("kind") == "document" or bool(data_source.get("textPreview") and not data_source.get("columns"))
    numeric_fields = _field_names(data_source, "number")
    all_fields = _field_names(data_source)
    metrics = model_payload.get("metrics") if model_payload else None
    dimensions = model_payload.get("dimensions") if model_payload else None
    charts = model_payload.get("charts") if model_payload else None

    if is_document:
        keywords = ((data_source.get("profile") or {}).get("keywords") or [])
        keyword_terms = [item.get("term") for item in keywords if item.get("term")]
        return {
            "request": request_text,
            "goal": (model_payload or {}).get("goal") or request_text,
            "metrics": metrics or "主题词频、证据匹配度、段落覆盖",
            "dimensions": dimensions or ", ".join(keyword_terms[:5]) or "章节、主题、风险点",
            "charts": charts or "主题词频柱状图、关键证据列表",
            "analysisMode": "document",
            "agents": [
                {"id": "source", "name": "数据接入 Agent", "output": "读取文档文本与基础画像"},
                {"id": "clean", "name": "文本清洗 Agent", "output": "清理空行并规范化段落"},
                {"id": "metric", "name": "主题提取 Agent", "output": "提取高频主题与关注点"},
                {"id": "chart", "name": "可视化 Agent", "output": "生成主题词频图"},
                {"id": "analysis", "name": "洞察分析 Agent", "output": "输出文档结论与证据"},
                {"id": "report", "name": "报告生成 Agent", "output": "生成文档分析报告"},
                {"id": "audit", "name": "结果校验 Agent", "output": "检查证据与范围说明"},
            ],
            "roleViews": {
                "business": "自然语言文档结论",
                "analyst": "主题词、证据段落与文本画像",
                "manager": "管理层决策简报",
            },
            "assets": ["文档画像", "主题词", "证据段落", "报告版本"],
            "schedule": "可对同类文档做定期复盘并输出主题变化",
        }

    return {
        "request": request_text,
        "goal": (model_payload or {}).get("goal") or request_text,
        "metrics": metrics or ", ".join(numeric_fields[:3]) or "自动识别核心数值指标",
        "dimensions": dimensions or ", ".join([field for field in all_fields if field not in numeric_fields][:3]) or "自动识别可比较维度",
        "charts": charts or "排名柱状图、占比环图、多指标雷达图",
        "analysisMode": "table",
        "agents": [
            {"id": "orchestrator", "name": "总调度 Agent", "output": "拆解需求、分派任务、汇总校验结论"},
            {"id": "source", "name": "数据接入 Agent", "output": "读取上传文件和数据源元信息"},
            {"id": "clean", "name": "数据清洗治理 Agent", "output": "识别缺失、重复、类型异常和口径风险"},
            {"id": "metric", "name": "指标计算 Agent", "output": "生成分组汇总、排名和对比指标"},
            {"id": "mining", "name": "挖掘建模 Agent", "output": "识别异常波动、集中度和潜在归因"},
            {"id": "chart", "name": "可视化 Agent", "output": "选择并生成对比图表"},
            {"id": "report", "name": "报告生成 Agent", "output": "输出业务版、管理版和技术版报告"},
            {"id": "audit", "name": "结果校验 Agent", "output": "交叉校验指标、图表和报告引用一致性"},
        ],
        "roleViews": {
            "business": "自然语言分析结论",
            "analyst": "SQL/Python 复用逻辑",
            "manager": "管理层决策简报",
        },
        "assets": ["分析模板", "清洗规则", "指标口径", "图表配置", "报告版本"],
        "schedule": "可接入定时巡检任务，按固定周期复跑同一分析目标并生成异常预警",
    }


def create_app(
    api_key: str,
    base_url: str,
    model: str,
    repository: TaskRepository | None = None,
    chat_client: AgnesChatClient | None = None,
    repository_path: str | None = None,
    database_url: str | None = None,
    admin_api_key: str = "",
    seed_demo_data: bool = True,
    model_timeout_seconds: float = 60.0,
    enable_audit_events: bool = True,
    cors_origins: list[str] | None = None,
    artifact_dir: str = "workspace-artifacts",
    max_request_body_bytes: int = 10_000_000,
    rate_limit_per_minute: int = 120,
    default_user_id: str = "local-user",
    default_tenant_id: str = "local-tenant",
    auth_secret: str = "local-dev-only-change-me",
    auth_token_minutes: int = 720,
    auth_required: bool = False,
    workflow_async: bool = False,
    workflow_step_delay_ms: int = 0,
    metrics_enabled: bool = True,
    queue_backend: str = "memory",
    redis_url: str | None = None,
    workflow_queue=None,
    error_tracking_enabled: bool = True,
    sentry_dsn: str | None = None,
    sentry_environment: str = "local",
    sentry_traces_sample_rate: float = 0.0,
    alert_webhook_url: str | None = None,
    alert_min_level: str = "error",
    error_tracker=None,
    rate_limit_backend: str = "memory",
    rate_limiter=None,
    credentials_secret: str = "local-dev-only-change-me",
    production_mode: bool = False,
    enforce_production_security: bool = True,
    enable_hsts: bool = False,
    content_security_policy: str | None = None,
    artifact_backend: str = "local",
    s3_bucket: str | None = None,
    s3_endpoint_url: str | None = None,
    s3_access_key: str | None = None,
    s3_secret_key: str | None = None,
    s3_region: str = "us-east-1",
    s3_prefix: str = "workspace-artifacts",
    web_fetch_timeout_seconds: float = 20.0,
    web_fetch_max_bytes: int = 5_000_000,
    web_table_max_rows: int = 1000,
    web_text_max_chars: int = 200_000,
) -> FastAPI:
    security_report = validate_production_settings(
        auth_required=auth_required,
        auth_secret=auth_secret,
        admin_api_key=admin_api_key,
        credentials_secret=credentials_secret,
        cors_origins=cors_origins,
        seed_demo_data=seed_demo_data,
        env={"WORKSPACE_PRODUCTION_MODE": "true" if production_mode else "false"},
    )
    if security_report.production_mode and enforce_production_security and not security_report.ok:
        raise RuntimeError(
            "production security validation failed: " + "; ".join(security_report.errors)
        )
    for warning in security_report.warnings:
        logger.warning("security_warning %s", warning)

    app = FastAPI()
    # Rate limiting supports memory (single instance) or redis (shared multi-instance).
    workflow_runner = workflow_queue or create_workflow_queue(
        backend=queue_backend,
        redis_url=redis_url,
        step_delay_seconds=max(0.0, workflow_step_delay_ms / 1000.0),
        max_workers=2,
    )
    app.state.workflow_runner = workflow_runner
    metrics = get_metrics()
    app.state.metrics = metrics
    tracker = error_tracker or init_error_tracking(
        enabled=error_tracking_enabled,
        sentry_dsn=sentry_dsn,
        environment=sentry_environment,
        traces_sample_rate=sentry_traces_sample_rate,
        alert_webhook_url=alert_webhook_url,
        alert_min_level=alert_min_level,
    )
    app.state.error_tracker = tracker
    limiter = rate_limiter or create_rate_limiter(
        backend=rate_limit_backend,
        redis_url=redis_url,
    )
    app.state.rate_limiter = limiter
    app.state.security_report = security_report
    allowed_origins = cors_origins or (
        []
        if production_mode
        else [
            "http://127.0.0.1:4173",
            "http://localhost:4173",
            "http://127.0.0.1:5173",
            "http://localhost:5173",
            "http://127.0.0.1:5174",
            "http://localhost:5174",
            "http://127.0.0.1:5175",
            "http://localhost:5175",
        ]
    )
    open_paths = public_auth_paths()

    def _public_user(user: AuthUser | dict) -> dict:
        if isinstance(user, AuthUser):
            role = user.role
            return {
                "id": user.id,
                "username": user.username,
                "tenantId": user.tenant_id,
                "displayName": user.display_name,
                "role": role,
                "permissions": permissions_for_role(role),
            }
        role = user["role"]
        return {
            "id": user["id"],
            "username": user["username"],
            "tenantId": user["tenantId"],
            "displayName": user["displayName"],
            "role": role,
            "permissions": permissions_for_role(role),
        }

    def _issue_token_for_user(user: dict) -> dict:
        auth_user = AuthUser(
            id=user["id"],
            username=user["username"],
            tenant_id=user["tenantId"],
            display_name=user["displayName"],
            role=user["role"],
        )
        token = create_access_token(
            user=auth_user,
            secret=auth_secret,
            expires_minutes=auth_token_minutes,
        )
        return {
            "accessToken": token,
            "tokenType": "Bearer",
            "expiresInMinutes": auth_token_minutes,
            "user": _public_user(auth_user),
        }

    def _extract_bearer_token(request: Request) -> str | None:
        authorization = (request.headers.get("authorization") or "").strip()
        if authorization.lower().startswith("bearer "):
            token = authorization[7:].strip()
            return token or None
        query_token = (request.query_params.get("accessToken") or "").strip()
        return query_token or None

    def _resolve_scope(request: Request) -> tuple[str, str, AuthUser | None]:
        token = _extract_bearer_token(request)
        if token:
            payload = decode_access_token(token, auth_secret)
            auth_user = user_from_token_payload(payload)
            return auth_user.tenant_id, auth_user.id, auth_user

        if auth_required and request.url.path not in open_paths:
            raise AuthError("authentication required", "auth_required")

        tenant_id = (
            request.headers.get("x-workspace-tenant-id")
            or request.query_params.get("workspaceTenantId")
            or default_tenant_id
        ).strip() or default_tenant_id
        owner_id = (
            request.headers.get("x-workspace-user-id")
            or request.query_params.get("workspaceUserId")
            or default_user_id
        ).strip() or default_user_id
        return tenant_id, owner_id, None

    def _effective_role(request: Request) -> str:
        auth_user = getattr(request.state, "auth_user", None)
        if auth_user is not None:
            return auth_user.role
        if auth_required:
            return ROLE_VIEWER
        return ROLE_LOCAL

    def require_permission(request: Request, permission: str) -> None:
        role = _effective_role(request)
        if not has_permission(role, permission):
            raise HTTPException(status_code=403, detail=f"{permission} permission required")

    def require_admin_access(request: Request) -> None:
        """Accept either a valid admin API key or an authenticated admin role."""
        if admin_api_key:
            provided = (request.headers.get("x-admin-api-key") or "").strip()
            if provided == admin_api_key:
                return
            if provided:
                raise HTTPException(status_code=403, detail="admin api key required")
            auth_user = getattr(request.state, "auth_user", None)
            if auth_user is not None and has_permission(auth_user.role, PERMISSION_ADMIN):
                return
            raise HTTPException(status_code=403, detail="admin api key or admin role required")
        require_permission(request, PERMISSION_ADMIN)

    @app.exception_handler(HTTPException)
    async def http_exception_handler(request: Request, exc: HTTPException):
        request_id = getattr(request.state, "request_id", str(uuid4()))
        detail = exc.detail if isinstance(exc.detail, str) else "request failed"
        return _error_response(
            request_id=request_id,
            status_code=exc.status_code,
            code=f"http_{exc.status_code}",
            message=detail,
            enable_hsts=enable_hsts,
            csp=content_security_policy,
        )

    @app.exception_handler(RequestValidationError)
    async def request_validation_handler(request: Request, exc: RequestValidationError):
        request_id = getattr(request.state, "request_id", str(uuid4()))
        return _error_response(
            request_id=request_id,
            status_code=422,
            code="request_validation_failed",
            message="request validation failed",
            details={"errors": exc.errors()},
            enable_hsts=enable_hsts,
            csp=content_security_policy,
        )

    @app.middleware("http")
    async def request_log_middleware(request: Request, call_next):
        request_id = request.headers.get("x-request-id") or str(uuid4())
        request.state.request_id = request_id
        request.state.auth_user = None
        try:
            tenant_id, owner_id, auth_user = _resolve_scope(request)
        except AuthError as error:
            status_code = 401
            metrics.observe_auth_failure()
            metrics.observe_request(
                method=request.method,
                path=request.url.path,
                status_code=status_code,
                duration_ms=0.0,
            )
            return _error_response(
                request_id=request_id,
                status_code=status_code,
                code=error.code,
                message=error.message,
            enable_hsts=enable_hsts,
            csp=content_security_policy,
        )

        request.state.auth_user = auth_user
        scope_token = set_workspace_scope(tenant_id, owner_id)
        try:
            started_at = time.perf_counter()
            client_host = request.client.host if request.client else "unknown"
            limit_key = rate_limit_key(
                client_host=client_host,
                tenant_id=tenant_id,
                user_id=owner_id if auth_user is not None else None,
            )
            decision = limiter.allow(limit_key, limit=rate_limit_per_minute, window_seconds=60)
            if rate_limit_per_minute > 0 and not decision.allowed:
                metrics.observe_rate_limited()
                metrics.observe_request(
                    method=request.method,
                    path=request.url.path,
                    status_code=429,
                    duration_ms=round((time.perf_counter() - started_at) * 1000, 2),
                )
                return _error_response(
                    request_id=request_id,
                    status_code=429,
                    code="rate_limit_exceeded",
                    message="too many requests",
                    details={
                        "retryAfterSeconds": decision.retry_after_seconds,
                        "limit": decision.limit,
                        "backend": decision.backend,
                    },
                    extra_headers={
                        "retry-after": str(decision.retry_after_seconds),
                        "x-ratelimit-limit": str(decision.limit),
                        "x-ratelimit-remaining": "0",
                        "x-ratelimit-backend": decision.backend,
                    },
            enable_hsts=enable_hsts,
            csp=content_security_policy,
        )
            content_length = request.headers.get("content-length")
            if content_length and int(content_length) > max_request_body_bytes:
                metrics.observe_request(
                    method=request.method,
                    path=request.url.path,
                    status_code=413,
                    duration_ms=round((time.perf_counter() - started_at) * 1000, 2),
                )
                return _error_response(
                    request_id=request_id,
                    status_code=413,
                    code="request_too_large",
                    message="request body too large",
            enable_hsts=enable_hsts,
            csp=content_security_policy,
        )
            if request.url.path not in open_paths and is_write_method(request.method):
                role = auth_user.role if auth_user else (ROLE_VIEWER if auth_required else ROLE_LOCAL)
                if not has_permission(role, PERMISSION_WRITE):
                    metrics.observe_request(
                        method=request.method,
                        path=request.url.path,
                        status_code=403,
                        duration_ms=round((time.perf_counter() - started_at) * 1000, 2),
                    )
                    return _error_response(
                        request_id=request_id,
                        status_code=403,
                        code="permission_denied",
                        message="write permission required",
                        details={"role": role, "required": PERMISSION_WRITE},
            enable_hsts=enable_hsts,
            csp=content_security_policy,
        )
            try:
                response = await call_next(request)
            except Exception as error:
                duration_ms = round((time.perf_counter() - started_at) * 1000, 2)
                metrics.observe_request(
                    method=request.method,
                    path=request.url.path,
                    status_code=500,
                    duration_ms=duration_ms,
                )
                tracker.capture_exception(
                    error,
                    source="http.middleware",
                    context={
                        "requestId": request_id,
                        "method": request.method,
                        "path": request.url.path,
                        "tenantId": tenant_id,
                        "userId": owner_id,
                        "durationMs": duration_ms,
                    },
                )
                logger.exception(
                    "request_failed method=%s path=%s duration_ms=%s request_id=%s",
                    request.method,
                    request.url.path,
                    duration_ms,
                    request_id,
                    extra={
                        "request_id": request_id,
                        "method": request.method,
                        "path": request.url.path,
                        "duration_ms": duration_ms,
                        "status": 500,
                        "tenant_id": tenant_id,
                        "user_id": owner_id,
                    },
                )
                return _error_response(
                    request_id=request_id,
                    status_code=500,
                    code="internal_server_error",
                    message="internal server error",
            enable_hsts=enable_hsts,
            csp=content_security_policy,
        )

            duration_ms = round((time.perf_counter() - started_at) * 1000, 2)
            if metrics_enabled and request.url.path not in {"/metrics", "/api/metrics"}:
                metrics.observe_request(
                    method=request.method,
                    path=request.url.path,
                    status_code=response.status_code,
                    duration_ms=duration_ms,
                )
            response.headers["x-request-id"] = request_id
            response.headers["x-workspace-tenant-id"] = tenant_id
            response.headers["x-workspace-user-id"] = owner_id
            response.headers["x-workspace-role"] = (
                auth_user.role if auth_user else (ROLE_VIEWER if auth_required else ROLE_LOCAL)
            )
            response.headers["x-ratelimit-limit"] = str(decision.limit)
            response.headers["x-ratelimit-remaining"] = str(decision.remaining)
            response.headers["x-ratelimit-backend"] = decision.backend
            _add_security_headers(response, enable_hsts=enable_hsts, csp=content_security_policy)
            log_method = logger.warning if response.status_code >= 500 else logger.info
            log_method(
                "request method=%s path=%s status=%s duration_ms=%s request_id=%s tenant_id=%s user_id=%s",
                request.method,
                request.url.path,
                response.status_code,
                duration_ms,
                request_id,
                tenant_id,
                owner_id,
                extra={
                    "request_id": request_id,
                    "method": request.method,
                    "path": request.url.path,
                    "status": response.status_code,
                    "duration_ms": duration_ms,
                    "tenant_id": tenant_id,
                    "user_id": owner_id,
                    "role": auth_user.role if auth_user else (ROLE_VIEWER if auth_required else ROLE_LOCAL),
                },
            )
            return response
        finally:
            reset_workspace_scope(scope_token)

    app.add_middleware(
        CORSMiddleware,
        allow_origins=allowed_origins,
        allow_credentials=True,
        allow_methods=["*"],
        allow_headers=["*"],
    )
    repo = repository or create_repository(
        db_path=repository_path,
        database_url=database_url,
        seed_demo_data=seed_demo_data,
    )
    client = chat_client or create_agnes_client(
        api_key=api_key,
        base_url=base_url,
        model=model,
        timeout_seconds=model_timeout_seconds,
    )
    store = create_artifact_store(
        backend=artifact_backend,
        artifact_dir=artifact_dir,
        s3_bucket=s3_bucket,
        s3_endpoint_url=s3_endpoint_url,
        s3_access_key=s3_access_key,
        s3_secret_key=s3_secret_key,
        s3_region=s3_region,
        s3_prefix=s3_prefix,
    )
    app.state.artifact_store = store
    app.state.artifact_backend = getattr(store, "backend", artifact_backend)

    @app.post("/api/auth/register")
    def register_user(payload: AuthRegisterRequest) -> dict:
        try:
            username = validate_username(payload.username)
            password = validate_password(payload.password)
        except AuthError as error:
            raise HTTPException(status_code=400, detail=error.message) from error

        tenant_id = (payload.tenantId or default_tenant_id).strip() or default_tenant_id
        display_name = (payload.displayName or username).strip() or username
        # First registered account becomes tenant admin; later accounts default to analyst.
        role = ROLE_ADMIN if repo.count_users() == 0 else ROLE_ANALYST
        try:
            user = repo.create_user(
                username=username,
                password_hash=hash_password(password),
                tenant_id=tenant_id,
                display_name=display_name,
                role=role,
            )
        except ValueError as error:
            raise HTTPException(status_code=409, detail=str(error)) from error

        scope_token = set_workspace_scope(user["tenantId"], user["id"])
        try:
            _audit(repo, "auth.register", user["id"], {"username": username, "tenantId": tenant_id, "role": role})
        finally:
            reset_workspace_scope(scope_token)
        return _issue_token_for_user(user)

    @app.post("/api/auth/login")
    def login_user(payload: AuthLoginRequest) -> dict:
        try:
            username = validate_username(payload.username)
            password = validate_password(payload.password)
        except AuthError as error:
            raise HTTPException(status_code=400, detail=error.message) from error

        user = repo.get_user_by_username(username)
        if not user or not verify_password(password, user["passwordHash"]):
            raise HTTPException(status_code=401, detail="invalid username or password")

        scope_token = set_workspace_scope(user["tenantId"], user["id"])
        try:
            _audit(repo, "auth.login", user["id"], {"username": username, "role": user["role"]})
        finally:
            reset_workspace_scope(scope_token)
        return _issue_token_for_user(user)

    @app.get("/api/auth/me")
    def current_user(request: Request) -> dict:
        auth_user = getattr(request.state, "auth_user", None)
        if auth_user is not None:
            return {
                "user": _public_user(auth_user),
                "authenticated": True,
                "permissions": permissions_for_role(auth_user.role),
            }

        if auth_required:
            raise HTTPException(status_code=401, detail="authentication required")

        return {
            "authenticated": False,
            "user": {
                "id": default_user_id,
                "username": default_user_id,
                "tenantId": default_tenant_id,
                "displayName": default_user_id,
                "role": ROLE_LOCAL,
                "permissions": permissions_for_role(ROLE_LOCAL),
            },
            "permissions": permissions_for_role(ROLE_LOCAL),
        }

    @app.get("/api/auth/users")
    def list_users(request: Request) -> dict:
        require_admin_access(request)
        auth_user = getattr(request.state, "auth_user", None)
        tenant_id = auth_user.tenant_id if auth_user else default_tenant_id
        users = repo.list_users(tenant_id=tenant_id)
        return {
            "items": [
                {
                    **item,
                    "permissions": permissions_for_role(item["role"]),
                }
                for item in users
            ]
        }

    @app.patch("/api/auth/users/{user_id}/role")
    def update_user_role(user_id: str, payload: RoleUpdateRequest, request: Request) -> dict:
        require_admin_access(request)
        try:
            role = normalize_role(payload.role)
        except AuthError as error:
            raise HTTPException(status_code=400, detail=error.message) from error

        target = repo.get_user_by_id(user_id)
        if not target:
            raise HTTPException(status_code=404, detail="user not found")

        auth_user = getattr(request.state, "auth_user", None)
        if auth_user and target["tenantId"] != auth_user.tenant_id:
            raise HTTPException(status_code=403, detail="cannot manage users outside your tenant")

        if auth_user and auth_user.id == user_id and role != ROLE_ADMIN:
            admins = [item for item in repo.list_users(tenant_id=target["tenantId"]) if item["role"] == ROLE_ADMIN]
            if len(admins) <= 1:
                raise HTTPException(status_code=400, detail="cannot demote the last admin")

        updated = repo.update_user_role(user_id, role)
        if not updated:
            raise HTTPException(status_code=404, detail="user not found")

        _audit(repo, "auth.role_update", user_id, {"role": role, "username": updated["username"]})
        return {
            "user": {
                "id": updated["id"],
                "username": updated["username"],
                "tenantId": updated["tenantId"],
                "displayName": updated["displayName"],
                "role": updated["role"],
                "permissions": permissions_for_role(updated["role"]),
            }
        }

    @app.get("/api/health")
    def health_check():
        model_status = {
            "configured": bool(api_key),
            "baseUrl": base_url,
            "name": model,
            "timeoutSeconds": model_timeout_seconds,
        }
        try:
            database = repo.ping()
        except Exception as error:
            return JSONResponse(
                status_code=503,
                content={
                    "status": "unavailable",
                    "database": {"ok": False, "error": str(error)},
                    "model": model_status,
                },
            )

        return {
            "status": "ok",
            "database": database,
            "model": model_status,
            "auth": {
                "required": auth_required,
                "tokenMinutes": auth_token_minutes,
                "roles": [ROLE_VIEWER, ROLE_ANALYST, ROLE_ADMIN],
            },
            "security": {
                "productionMode": production_mode,
                "hsts": enable_hsts,
                "enforceProductionSecurity": enforce_production_security,
                "ok": security_report.ok,
                "warnings": security_report.warnings,
            },
            "storage": {
                "backend": getattr(store, "backend", artifact_backend),
            },
            "observability": {
                "metricsEnabled": metrics_enabled,
                "metricsPath": "/metrics",
                "snapshot": metrics.snapshot() if metrics_enabled else {},
                "errorTracking": tracker.snapshot(),
            },
            "queue": {
                "backend": getattr(workflow_runner, "backend_name", queue_backend),
                "async": workflow_async,
            },
            "rateLimit": {
                "backend": getattr(limiter, "backend_name", rate_limit_backend),
                "perMinute": rate_limit_per_minute,
            },
        }

    @app.get("/metrics")
    def prometheus_metrics():
        if not metrics_enabled:
            raise HTTPException(status_code=404, detail="metrics disabled")
        from fastapi.responses import PlainTextResponse

        return PlainTextResponse(metrics.render_prometheus(), media_type="text/plain; version=0.0.4; charset=utf-8")

    @app.get("/api/metrics")
    def metrics_snapshot():
        if not metrics_enabled:
            raise HTTPException(status_code=404, detail="metrics disabled")
        return {
            "status": "ok",
            "metrics": metrics.snapshot(),
            "errorTracking": tracker.snapshot(),
        }

    @app.get("/api/errors/recent")
    def recent_errors(request: Request, limit: int = 20):
        require_admin_access(request)
        return {
            "items": tracker.recent_events(limit=limit),
            "summary": tracker.snapshot(),
        }

    @app.get("/api/health/live")
    def live_check():
        return {
            "status": "ok",
            "service": "workspace-api",
            "time": datetime.now(timezone.utc).isoformat(),
        }

    @app.get("/api/health/ready")
    def readiness_check():
        try:
            database = repo.ping()
        except Exception as error:
            return JSONResponse(
                status_code=503,
                content={
                    "status": "unavailable",
                    "checks": {
                        "database": {"ok": False, "error": str(error)},
                        "model": {"ok": bool(api_key), "configured": bool(api_key)},
                    },
                },
            )
        return {
            "status": "ok",
            "checks": {
                "database": {"ok": True, "detail": database},
                "model": {"ok": bool(api_key), "configured": bool(api_key), "name": model},
            },
        }

    @app.get("/api/active-task")
    def get_active_task() -> dict:
        tenant_id, owner_id = get_workspace_scope()
        job_view = workflow_runner.job_public_view(workflow_runner.get_job(tenant_id, owner_id))
        return attach_execution_meta(repo.get_active_task(), job_view, mode="async" if workflow_async else "sync")

    @app.patch("/api/active-task")
    def patch_active_task(update: ActiveTaskUpdate) -> dict:
        payload = update.model_dump(exclude_none=True)
        task = repo.update_active_task(payload)
        _audit(repo, "active_task.patch", task["id"], {"fields": sorted(payload.keys())})
        tenant_id, owner_id = get_workspace_scope()
        job_view = workflow_runner.job_public_view(workflow_runner.get_job(tenant_id, owner_id))
        return attach_execution_meta(task, job_view, mode="async" if workflow_async else "sync")

    @app.post("/api/active-task/reset")
    def reset_active_task() -> dict:
        task = repo.reset_active_task()
        _audit(repo, "active_task.reset", task["id"])
        return task

    @app.post("/api/active-task/from-replay")
    def create_task_from_replay() -> dict:
        task = repo.create_task_from_replay()
        _audit(repo, "active_task.from_replay", task["id"])
        return task

    @app.post("/api/history/save-active-task")
    def save_active_task_to_history() -> dict:
        task = repo.save_active_task_to_history()
        _audit(repo, "history.save_active_task", task["id"])
        return task

    @app.get("/api/history/tasks")
    def list_history_tasks() -> list[dict]:
        return repo.list_recent_tasks()

    @app.get("/api/history/tasks/{task_id}")
    def get_history_task_detail(task_id: str) -> dict | None:
        return repo.get_task_detail(task_id)

    @app.post("/api/history/tasks/{task_id}/activate")
    def activate_history_task(task_id: str) -> dict:
        task = repo.load_history_task_into_active(task_id)
        if not task:
            raise HTTPException(status_code=404, detail="history task not found")
        _audit(repo, "history.activate", task_id)
        return task

    @app.delete("/api/history/tasks/{task_id}")
    def delete_history_task(task_id: str, request: Request) -> dict:
        require_admin_access(request)
        task = repo.delete_history_task(task_id)
        if not task:
            raise HTTPException(status_code=404, detail="history task not found")
        for artifact in repo.delete_artifacts_for_task(task_id):
            storage_key = artifact.get("storageKey")
            if storage_key:
                store.delete(storage_key)
        _audit(repo, "history.delete", task_id)
        return {"deleted": True, "id": task_id}

    @app.post("/api/templates/{template_id}/activate")
    def activate_template(template_id: str) -> dict:
        task = repo.create_task_from_template(template_id)
        if not task:
            raise HTTPException(status_code=404, detail="template not found")
        _audit(repo, "template.activate", template_id)
        return task

    @app.get("/api/templates")
    def list_templates() -> list[dict]:
        return repo.list_templates()

    @app.get("/api/datasets")
    def list_datasets() -> list[dict]:
        return repo.list_datasets()

    @app.get("/api/connectors/types")
    def get_connector_types() -> dict:
        return {"items": list_connector_types()}

    @app.get("/api/connections")
    def list_connections() -> dict:
        return {"items": [public_connection_view(item) for item in repo.list_connections()]}

    @app.post("/api/connections")
    def create_connection(payload: ConnectionUpsertRequest) -> dict:
        connector_type = payload.connectorType.strip().lower()
        if connector_type not in {item["type"] for item in list_connector_types()}:
            raise HTTPException(status_code=400, detail="unsupported connector type")
        name = payload.name.strip()
        if not name:
            raise HTTPException(status_code=400, detail="name is required")
        credentials = payload.credentials or {}
        connection = {
            "name": name,
            "connectorType": connector_type,
            "config": payload.config or {},
            "credentialsEncrypted": encrypt_credentials(credentials, credentials_secret) if credentials else "",
        }
        saved = repo.save_connection(connection)
        _audit(repo, "connection.create", saved["id"], {"connectorType": connector_type, "name": name})
        return public_connection_view(saved)

    @app.patch("/api/connections/{connection_id}")
    def update_connection(connection_id: str, payload: ConnectionUpsertRequest) -> dict:
        existing = repo.get_connection(connection_id)
        if not existing:
            raise HTTPException(status_code=404, detail="connection not found")
        connector_type = (payload.connectorType or existing.get("connectorType") or "").strip().lower()
        if connector_type not in {item["type"] for item in list_connector_types()}:
            raise HTTPException(status_code=400, detail="unsupported connector type")
        name = payload.name.strip() or existing.get("name") or connection_id
        next_connection = {
            **existing,
            "name": name,
            "connectorType": connector_type,
            "config": payload.config if payload.config is not None else existing.get("config") or {},
        }
        if payload.credentials is not None:
            next_connection["credentialsEncrypted"] = (
                encrypt_credentials(payload.credentials, credentials_secret) if payload.credentials else ""
            )
        saved = repo.save_connection(next_connection)
        _audit(repo, "connection.update", connection_id, {"connectorType": connector_type, "name": name})
        return public_connection_view(saved)

    @app.delete("/api/connections/{connection_id}")
    def delete_connection(connection_id: str, request: Request) -> dict:
        require_admin_access(request)
        deleted = repo.delete_connection(connection_id)
        if not deleted:
            raise HTTPException(status_code=404, detail="connection not found")
        _audit(repo, "connection.delete", connection_id)
        return {"deleted": True, "id": connection_id}

    @app.post("/api/connections/{connection_id}/test")
    def test_saved_connection(connection_id: str, payload: ConnectionTestRequest | None = None) -> dict:
        connection = repo.get_connection(connection_id)
        if not connection:
            raise HTTPException(status_code=404, detail="connection not found")
        body = payload or ConnectionTestRequest()
        connector_type = (body.connectorType or connection.get("connectorType") or "").strip().lower()
        config = body.config if body.config is not None else (connection.get("config") or {})
        if body.credentials is not None:
            credentials = body.credentials
        elif connection.get("credentialsEncrypted"):
            try:
                credentials = decrypt_credentials(connection["credentialsEncrypted"], credentials_secret)
            except ConnectorError as error:
                raise HTTPException(status_code=400, detail=error.message) from error
        else:
            credentials = {}
        try:
            result = test_connection(connector_type, config, credentials)
        except ConnectorError as error:
            connection["lastTestedAt"] = datetime.now(timezone.utc).isoformat()
            connection["lastTestStatus"] = "failed"
            repo.save_connection(connection)
            raise HTTPException(status_code=400, detail=error.message) from error
        connection["lastTestedAt"] = datetime.now(timezone.utc).isoformat()
        connection["lastTestStatus"] = "ok"
        repo.save_connection(connection)
        _audit(repo, "connection.test", connection_id, {"ok": True, "rowCount": result.get("rowCount")})
        return result

    @app.post("/api/connections/test")
    def test_ephemeral_connection(payload: ConnectionTestRequest) -> dict:
        connector_type = (payload.connectorType or "").strip().lower()
        if not connector_type:
            raise HTTPException(status_code=400, detail="connectorType is required")
        try:
            return test_connection(connector_type, payload.config or {}, payload.credentials or {})
        except ConnectorError as error:
            raise HTTPException(status_code=400, detail=error.message) from error

    @app.post("/api/connections/{connection_id}/import")
    def import_connection_dataset(connection_id: str, payload: ConnectionImportRequest | None = None) -> dict:
        connection = repo.get_connection(connection_id)
        if not connection:
            raise HTTPException(status_code=404, detail="connection not found")
        body = payload or ConnectionImportRequest()
        credentials = {}
        if connection.get("credentialsEncrypted"):
            try:
                credentials = decrypt_credentials(connection["credentialsEncrypted"], credentials_secret)
            except ConnectorError as error:
                raise HTTPException(status_code=400, detail=error.message) from error
        try:
            payload = fetch_connector_payload(
                connection.get("connectorType") or "",
                connection.get("config") or {},
                credentials,
            )
            data_source = build_connector_data_source(
                name=body.datasetName or connection.get("name") or connection_id,
                connector_type=str(connection.get("connectorType") or "connector"),
                payload=payload,
                connection_id=connection_id,
            )
        except ConnectorError as error:
            raise HTTPException(status_code=400, detail=error.message) from error
        data_source["uploadedAt"] = datetime.now(timezone.utc).isoformat()
        dataset = repo.add_uploaded_dataset(data_source)
        connection["lastImportAt"] = datetime.now(timezone.utc).isoformat()
        repo.save_connection(connection)
        _audit(
            repo,
            "connection.import",
            connection_id,
            {
                "datasetId": dataset.get("id"),
                "kind": dataset.get("kind"),
                "rowCount": dataset.get("rowCount"),
                "columnCount": dataset.get("columnCount"),
            },
        )
        return dataset

    @app.delete("/api/datasets/{dataset_id}")
    def delete_dataset(dataset_id: str, request: Request) -> dict:
        require_admin_access(request)
        if repo.active_task_uses_dataset(dataset_id):
            raise HTTPException(status_code=409, detail="active task is using this dataset")
        dataset = repo.delete_dataset(dataset_id)
        if not dataset:
            raise HTTPException(status_code=404, detail="dataset not found")
        storage_key = dataset.get("storageKey")
        if storage_key:
            store.delete(storage_key)
        _audit(repo, "dataset.delete", dataset_id)
        return {"deleted": True, "id": dataset_id}

    @app.get("/api/reports")
    def list_reports() -> list[dict]:
        return repo.list_reports()

    @app.get("/api/search")
    def search(query: str = "") -> list[dict]:
        return repo.search(query)

    @app.get("/api/notifications")
    def list_notifications() -> list[dict]:
        return repo.list_notifications()

    def persist_uploaded_data_source(file_name: str, mime_type: str | None, content: bytes) -> dict:
        try:
            checked = validate_upload(
                file_name=file_name,
                content_type=mime_type,
                size_bytes=len(content),
                max_bytes=max_request_body_bytes,
                content=content,
            )
        except UploadSecurityError as error:
            raise HTTPException(status_code=400, detail=error.message) from error
        encoded = base64.b64encode(content).decode("ascii")
        try:
            parsed = ingest_data_source(checked["fileName"], checked["contentType"], encoded)
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error)) from error
        stored = store.save("uploads", checked["fileName"], content)
        parsed["storageKey"] = stored["storageKey"]
        parsed["originalMimeType"] = checked["contentType"]
        parsed["originalFileSize"] = len(content)
        parsed["uploadedAt"] = datetime.now(timezone.utc).isoformat()
        parsed["storageBackend"] = stored.get("backend") or artifact_backend
        dataset = repo.add_uploaded_dataset(parsed)
        _audit(
            repo,
            "data_source.ingest",
            dataset["id"],
            {
                "kind": parsed["kind"],
                "formatLabel": parsed["formatLabel"],
                "rowCount": parsed["rowCount"],
                "storageBackend": parsed.get("storageBackend"),
            },
        )
        return dataset

    @app.post("/api/data-source/upload")
    async def upload_data_source(file: UploadFile = File(...)) -> dict:
        content = await file.read()
        if len(content) > max_request_body_bytes:
            raise HTTPException(status_code=413, detail="uploaded file exceeds request size limit")
        return persist_uploaded_data_source(file.filename or "upload", file.content_type, content)

    @app.post("/api/data-source/ingest")
    def ingest_uploaded_data_source(payload: DataSourceIngestRequest) -> dict:
        try:
            content = base64.b64decode(payload.contentBase64, validate=True)
        except ValueError as error:
            raise HTTPException(status_code=400, detail="invalid file payload") from error
        if len(content) > max_request_body_bytes:
            raise HTTPException(status_code=413, detail="request body too large")
        try:
            return persist_uploaded_data_source(payload.fileName, payload.mimeType, content)
        except HTTPException as error:
            raise error

    @app.get("/api/datasets/{dataset_id}/download")
    def download_dataset(dataset_id: str):
        dataset = repo.get_dataset(dataset_id)
        if not dataset or not dataset.get("storageKey"):
            raise HTTPException(status_code=404, detail="dataset file not found")
        try:
            path = store.resolve(dataset["storageKey"])
        except (OSError, ValueError) as error:
            raise HTTPException(status_code=404, detail="dataset file not found") from error
        return FileResponse(path, media_type=dataset.get("originalMimeType") or "application/octet-stream", filename=dataset["name"])

    @app.get("/api/tasks/active")
    def get_canonical_active_task() -> dict:
        tenant_id, owner_id = get_workspace_scope()
        job_view = workflow_runner.job_public_view(workflow_runner.get_job(tenant_id, owner_id))
        return attach_execution_meta(repo.get_active_task(), job_view, mode="async" if workflow_async else "sync")

    @app.post("/api/tasks/active/dataset/{dataset_id}")
    def attach_dataset_to_active_task(dataset_id: str) -> dict:
        dataset = repo.get_dataset(dataset_id)
        if not dataset:
            raise HTTPException(status_code=404, detail="dataset not found")
        task = repo.update_active_task({"dataSource": dataset, "status": "等待需求解析"})
        _audit(repo, "task.dataset.attach", task["id"], {"datasetId": dataset_id})
        return task

    @app.delete("/api/tasks/active/report")
    def delete_active_report(request: Request) -> dict:
        require_admin_access(request)
        for artifact in repo.delete_artifacts_for_task("workspace-active-session"):
            storage_key = artifact.get("storageKey")
            if storage_key:
                store.delete(storage_key)
        task = repo.reset_active_task()
        _audit(repo, "active_task.report_delete", task["id"])
        return {"deleted": True, "id": task["id"]}

    @app.post("/api/tasks/active/correction")
    def apply_task_correction(payload: TaskCorrectionRequest) -> dict:
        task = repo.get_active_task()
        snapshot = task["snapshot"]
        nodes = [{**node, "status": "success" if node["id"] in {"source", "clean"} else "pending"} for node in snapshot["nodes"]]
        snapshot["nodes"] = nodes
        snapshot["phase"] = "idle"
        snapshot["logMessage"] = "已记录人工修正，指标及后续节点等待重新执行"
        snapshot["context"] = {**snapshot["context"], "latestLog": snapshot["logMessage"]}
        task = repo.update_active_task({"summary": payload.summary, "snapshot": snapshot, "status": "等待重新执行", "artifacts": [item for item in task.get("artifacts", []) if item.get("nodeId") in {"source", "clean"}]})
        _audit(repo, "task.correction", task["id"])
        return task

    @app.post("/api/tasks/active/exports")
    def export_active_task(payload: ExportRequest) -> dict:
        task = repo.get_active_task()
        audit = next((item for item in reversed(task.get("artifacts", [])) if item.get("type") == "audit"), None)
        if not audit:
            raise HTTPException(status_code=409, detail="complete workflow audit before exporting")
        report = payload.report or next((item.get("content", {}).get("markdown") for item in reversed(task.get("artifacts", [])) if item.get("type") == "report_markdown"), task.get("summary", ""))
        export_format = payload.format.lower()
        safe_name = "".join(character if character not in '\\/:*?\"<>|' else "_" for character in task.get("name", "analysis-report"))
        if export_format in {"markdown", "md"}:
            file_name, mime_type, content = f"{safe_name}.md", "text/markdown; charset=utf-8", report.encode("utf-8")
        elif export_format == "html":
            file_name, mime_type, content = f"{safe_name}.html", "text/html; charset=utf-8", f"<!doctype html><meta charset='utf-8'><title>{escape(task.get('name', '分析报告'))}</title><pre>{escape(report)}</pre>".encode("utf-8")
        elif export_format in {"word", "docx"}:
            from docx import Document
            document = Document()
            document.add_heading(task.get("name", "分析报告"), 0)
            for line in report.splitlines():
                document.add_paragraph(line)
            temp = store.root / "exports" / f"{uuid4()}.docx"
            temp.parent.mkdir(parents=True, exist_ok=True)
            document.save(temp)
            content = temp.read_bytes()
            temp.unlink(missing_ok=True)
            file_name, mime_type = f"{safe_name}.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif export_format == "pdf":
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.pdfgen import canvas
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
            temp = store.root / "exports" / f"{uuid4()}.pdf"
            temp.parent.mkdir(parents=True, exist_ok=True)
            pdf = canvas.Canvas(str(temp), pagesize=A4)
            pdf.setFont("STSong-Light", 11)
            y = 800
            for line in report.splitlines() or ["分析报告"]:
                pdf.drawString(48, y, line[:90])
                y -= 18
                if y < 48:
                    pdf.showPage()
                    pdf.setFont("STSong-Light", 11)
                    y = 800
            pdf.save()
            content = temp.read_bytes()
            temp.unlink(missing_ok=True)
            file_name, mime_type = f"{safe_name}.pdf", "application/pdf"
        else:
            raise HTTPException(status_code=400, detail="unsupported export format")
        stored = store.save("exports", file_name, content)
        artifact = {**stored, "taskId": task["id"], "type": "export", "format": export_format, "mimeType": mime_type, "createdAt": datetime.now(timezone.utc).isoformat()}
        repo.add_artifact(artifact)
        exports = [{"format": export_format.upper(), "status": f"{export_format.upper()} 导出完成", "artifactId": artifact["id"]}, *task.get("exports", [])]
        repo.update_active_task({"exports": exports})
        _audit(repo, "task.export", task["id"], {"artifactId": artifact["id"], "format": export_format})
        return {**artifact, "downloadUrl": f"/api/artifacts/{artifact['id']}/download"}

    @app.get("/api/artifacts/{artifact_id}/download")
    def download_artifact(artifact_id: str):
        artifact = repo.get_artifact(artifact_id)
        if not artifact:
            raise HTTPException(status_code=404, detail="artifact not found")
        try:
            path = store.resolve(artifact["storageKey"])
        except (OSError, ValueError) as error:
            raise HTTPException(status_code=404, detail="artifact file not found") from error
        return FileResponse(path, media_type=artifact.get("mimeType") or "application/octet-stream", filename=artifact.get("fileName") or "artifact")

    @app.post("/api/templates/save-active-task")
    def save_active_task_as_template(payload: TemplateDraftRequest, request: Request) -> dict:
        require_admin_access(request)
        template = repo.save_active_task_as_template(payload.model_dump(exclude_none=True))
        _audit(repo, "template.save_active_task", template["id"])
        return template

    @app.get("/api/audit/events")
    def list_audit_events(request: Request, limit: int = 50, offset: int = 0) -> dict:
        require_admin_access(request)
        if not enable_audit_events:
            raise HTTPException(status_code=404, detail="Not found")
        safe_limit = max(1, min(limit, 200))
        safe_offset = max(0, offset)
        total = repo.count_audit_events()
        items = repo.list_audit_events(limit=safe_limit, offset=safe_offset)
        return {
            "items": items,
            "pagination": {
                "limit": safe_limit,
                "offset": safe_offset,
                "total": total,
                "hasMore": safe_offset + len(items) < total,
            },
        }

    @app.post("/api/intent/parse")
    async def parse_intent(payload: IntentRequest):
        request_text = payload.request.strip()
        if not request_text:
            raise HTTPException(status_code=400, detail="request is required")

        prompt = "\n".join(
            [
                "Convert the analytics request below into JSON only.",
                "Return only JSON.",
                "Use keys: request, goal, metrics, dimensions, charts.",
                f"Request: {request_text}{_data_source_context(payload.dataSource)}",
            ]
        )

        try:
            content = await client.complete(prompt)
            model_payload = json.loads(content)
            if not isinstance(model_payload, dict):
                model_payload = {}
            return _build_agentic_intent(request_text, payload.dataSource, model_payload)
        except Exception:
            return _build_agentic_intent(request_text, payload.dataSource)

    @app.post("/api/workflow/action")
    def workflow_action(payload: WorkflowActionRequest) -> dict:
        action = (payload.action or "").strip()
        if not action:
            raise HTTPException(status_code=400, detail="action is required")

        tenant_id, owner_id = get_workspace_scope()
        control_actions = {"pause", "stop"}
        run_actions = {"start", "resume", "retry", "retry_current"}

        if action in control_actions:
            if action == "stop":
                workflow_runner.request_cancel(tenant_id, owner_id)
            task = repo.apply_workflow_action(action, chat_client=None, run_to_completion=False)
            _audit(repo, "workflow.action", task["id"], {"action": action, "phase": task.get("snapshot", {}).get("phase"), "mode": "control"})
            job_view = workflow_runner.job_public_view(workflow_runner.get_job(tenant_id, owner_id))
            return attach_execution_meta(task, job_view, mode="async" if workflow_async else "sync")

        if action == "fail":
            task = repo.apply_workflow_action(action, chat_client=None, run_to_completion=False)
            _audit(repo, "workflow.action", task["id"], {"action": action, "phase": task.get("snapshot", {}).get("phase")})
            job_view = workflow_runner.job_public_view(workflow_runner.get_job(tenant_id, owner_id))
            return attach_execution_meta(task, job_view, mode="async" if workflow_async else "sync")

        if action not in run_actions:
            raise HTTPException(status_code=400, detail=f"unsupported action: {action}")

        if workflow_runner.is_active(tenant_id, owner_id):
            raise HTTPException(status_code=409, detail="workflow job already running")

        # Transition into running state immediately so clients see progress on first poll.
        task = repo.apply_workflow_action(action, chat_client=None, run_to_completion=False)

        if not workflow_async:
            task = repo.run_workflow_until_idle(client)
            terminal = "succeeded" if task.get("snapshot", {}).get("phase") != "failed" else "failed"
            task = repo.set_execution_meta(
                {
                    "mode": "sync",
                    "backend": getattr(workflow_runner, "backend_name", "memory"),
                    "status": terminal,
                    "action": action,
                    "finishedAt": datetime.now(timezone.utc).isoformat(),
                }
            )
            metrics.observe_workflow_job(terminal, steps=len(task.get("artifacts") or []))
            _audit(repo, "workflow.action", task["id"], {"action": action, "phase": task.get("snapshot", {}).get("phase"), "mode": "sync"})
            return attach_execution_meta(task, task.get("execution"), mode="sync")

        queue_backend_name = getattr(workflow_runner, "backend_name", "memory")

        def run_step() -> dict:
            return repo.advance_workflow(client)

        def should_continue(current_task: dict) -> bool:
            return current_task.get("snapshot", {}).get("phase") == "running"

        def on_complete(job, final_task):
            repo.set_execution_meta(
                {
                    "mode": "async",
                    "backend": queue_backend_name,
                    "jobId": job.job_id,
                    "status": job.status,
                    "action": job.action,
                    "submittedAt": job.submitted_at,
                    "startedAt": job.started_at,
                    "finishedAt": job.finished_at,
                    "lastError": job.last_error,
                    "steps": job.steps,
                }
            )
            metrics.observe_workflow_job(job.status, steps=job.steps)
            if job.status == "failed":
                tracker.capture_message(
                    job.last_error or "workflow job failed",
                    source="workflow.job",
                    level="error",
                    context={
                        "jobId": job.job_id,
                        "action": job.action,
                        "tenantId": job.tenant_id,
                        "ownerId": job.owner_id,
                        "steps": job.steps,
                        "backend": queue_backend_name,
                    },
                )
            _audit(
                repo,
                "workflow.job.complete",
                (final_task or {}).get("id", "workspace-active-session"),
                {"jobId": job.job_id, "status": job.status, "steps": job.steps, "error": job.last_error, "backend": queue_backend_name},
            )

        if queue_backend_name == "redis":
            # API only enqueues; external worker process advances nodes.
            job = workflow_runner.submit(
                tenant_id=tenant_id,
                owner_id=owner_id,
                action=action,
            )
        else:
            job = workflow_runner.submit(
                tenant_id=tenant_id,
                owner_id=owner_id,
                action=action,
                run_step=run_step,
                should_continue=should_continue,
                on_complete=on_complete,
            )
        job_view = workflow_runner.job_public_view(job)
        task = repo.set_execution_meta(
            {
                "mode": "async",
                "backend": queue_backend_name,
                **(job_view or {}),
            }
        )
        _audit(
            repo,
            "workflow.action",
            task["id"],
            {
                "action": action,
                "phase": task.get("snapshot", {}).get("phase"),
                "mode": "async",
                "backend": queue_backend_name,
                "jobId": job.job_id,
            },
        )
        return attach_execution_meta(task, job_view, mode="async")

    @app.get("/api/workflow/jobs/current")
    def current_workflow_job() -> dict:
        tenant_id, owner_id = get_workspace_scope()
        job_view = workflow_runner.job_public_view(workflow_runner.get_job(tenant_id, owner_id))
        task = attach_execution_meta(repo.get_active_task(), job_view, mode="async" if workflow_async else "sync")
        return {
            "job": job_view,
            "task": task,
        }

    @app.get("/api/workflow/jobs/stream")
    async def workflow_job_stream(request: Request):
        """Server-Sent Events stream for current workflow job progress."""
        import asyncio
        from fastapi.responses import StreamingResponse

        tenant_id, owner_id = get_workspace_scope()

        async def event_generator():
            last_payload = None
            # Keep connection for up to ~10 minutes by default.
            for _ in range(600):
                if await request.is_disconnected():
                    break
                job_view = workflow_runner.job_public_view(workflow_runner.get_job(tenant_id, owner_id))
                task = attach_execution_meta(
                    repo.get_active_task(),
                    job_view,
                    mode="async" if workflow_async else "sync",
                )
                payload = {
                    "job": job_view,
                    "task": {
                        "id": task.get("id"),
                        "status": task.get("status"),
                        "snapshot": task.get("snapshot"),
                        "execution": task.get("execution"),
                    },
                    "ts": datetime.now(timezone.utc).isoformat(),
                }
                serialized = json.dumps(payload, ensure_ascii=False)
                if serialized != last_payload:
                    yield f"event: workflow\ndata: {serialized}\n\n"
                    last_payload = serialized
                status = (job_view or {}).get("status") or (task.get("execution") or {}).get("status")
                phase = ((task.get("snapshot") or {}).get("phase"))
                if status in {"succeeded", "failed", "cancelled"} or phase in {
                    "completed",
                    "needs_review",
                    "failed",
                    "stopped",
                }:
                    yield f"event: done\ndata: {serialized}\n\n"
                    break
                await asyncio.sleep(1.0)

        headers = {
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        }
        return StreamingResponse(event_generator(), media_type="text/event-stream", headers=headers)

    @app.post("/api/workflow/advance")
    def workflow_advance() -> dict:
        # Compatibility endpoint: prefer POST /api/workflow/action with start/resume.
        task = repo.advance_workflow(client)
        _audit(repo, "workflow.advance", task["id"])
        tenant_id, owner_id = get_workspace_scope()
        job_view = workflow_runner.job_public_view(workflow_runner.get_job(tenant_id, owner_id))
        return attach_execution_meta(task, job_view, mode="async" if workflow_async else "sync")

    return app
