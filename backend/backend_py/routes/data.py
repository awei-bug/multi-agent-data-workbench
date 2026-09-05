from __future__ import annotations

import base64
import json
from datetime import datetime, timezone
from html import escape
from pathlib import Path
from uuid import uuid4

from fastapi import APIRouter, FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse

from backend_py.api_helpers import audit, build_agentic_intent, data_source_context
from backend_py.app_context import AppContext
from backend_py.compliance import (
    apply_html_watermark,
    apply_text_watermark,
    build_watermark_lines,
    mask_text,
)
from backend_py.connectors import (
    ConnectorError,
    build_connector_data_source,
    decrypt_credentials,
    encrypt_credentials,
    fetch_connector_payload,
    list_connector_types,
    public_connection_view,
    test_connection,
)
from backend_py.data_source_ingest import ingest_data_source
from backend_py.repository import get_workspace_scope
from backend_py.schemas import (
    ActiveTaskUpdate,
    ConnectionImportRequest,
    ConnectionTestRequest,
    ConnectionUpsertRequest,
    DataSourceIngestRequest,
    ExportRequest,
    IntentRequest,
    SmartScrapeCommitRequest,
    SmartScrapeRequest,
    TaskCorrectionRequest,
    TemplateDraftRequest,
)
from backend_py.smart_scrape import smart_scrape_to_dataset
from backend_py.upload_security import UploadSecurityError, validate_upload
from backend_py.workflow_runner import attach_execution_meta
import asyncio
import inspect


def register_data_routes(app: FastAPI, ctx: AppContext) -> None:
    router = APIRouter(tags=["data"])

    @router.get("/api/active-task")
    def get_active_task() -> dict:
        tenant_id, owner_id = get_workspace_scope()
        job_view = ctx.workflow_runner.job_public_view(ctx.workflow_runner.get_job(tenant_id, owner_id))
        return attach_execution_meta(ctx.repo.get_active_task(), job_view, mode="async" if ctx.workflow_async else "sync")

    @router.patch("/api/active-task")
    def patch_active_task(update: ActiveTaskUpdate) -> dict:
        payload = update.model_dump(exclude_none=True)
        task = ctx.repo.update_active_task(payload)
        audit(ctx.repo, "active_task.patch", task["id"], {"fields": sorted(payload.keys())})
        tenant_id, owner_id = get_workspace_scope()
        job_view = ctx.workflow_runner.job_public_view(ctx.workflow_runner.get_job(tenant_id, owner_id))
        return attach_execution_meta(task, job_view, mode="async" if ctx.workflow_async else "sync")

    @router.post("/api/active-task/reset")
    def reset_active_task() -> dict:
        task = ctx.repo.reset_active_task()
        audit(ctx.repo, "active_task.reset", task["id"])
        return task

    @router.post("/api/active-task/from-replay")
    def create_task_from_replay() -> dict:
        task = ctx.repo.create_task_from_replay()
        audit(ctx.repo, "active_task.from_replay", task["id"])
        return task

    @router.post("/api/history/save-active-task")
    def save_active_task_to_history() -> dict:
        task = ctx.repo.save_active_task_to_history()
        audit(ctx.repo, "history.save_active_task", task["id"])
        return task

    @router.get("/api/history/tasks")
    def list_history_tasks() -> list[dict]:
        return ctx.repo.list_recent_tasks()

    @router.get("/api/history/tasks/{task_id}")
    def get_history_task_detail(task_id: str) -> dict | None:
        return ctx.repo.get_task_detail(task_id)

    @router.post("/api/history/tasks/{task_id}/activate")
    def activate_history_task(task_id: str) -> dict:
        task = ctx.repo.load_history_task_into_active(task_id)
        if not task:
            raise HTTPException(status_code=404, detail="history task not found")
        audit(ctx.repo, "history.activate", task_id)
        return task

    @router.delete("/api/history/tasks/{task_id}")
    def delete_history_task(task_id: str, request: Request) -> dict:
        ctx.require_admin_access(request)
        task = ctx.repo.delete_history_task(task_id)
        if not task:
            raise HTTPException(status_code=404, detail="history task not found")
        tenant_id, _owner_id = get_workspace_scope()
        for artifact in ctx.repo.delete_artifacts_for_task(task_id):
            storage_key = artifact.get("storageKey")
            if storage_key:
                try:
                    ctx.store.delete(storage_key, tenant_id=tenant_id)
                except Exception:
                    pass
        audit(ctx.repo, "history.delete", task_id)
        return {"deleted": True, "id": task_id}

    @router.post("/api/templates/{template_id}/activate")
    def activate_template(template_id: str) -> dict:
        task = ctx.repo.create_task_from_template(template_id)
        if not task:
            raise HTTPException(status_code=404, detail="template not found")
        audit(ctx.repo, "template.activate", template_id)
        return task

    @router.get("/api/templates")
    def list_templates() -> list[dict]:
        return ctx.repo.list_templates()

    @router.get("/api/datasets")
    def list_datasets() -> list[dict]:
        return ctx.repo.list_datasets()

    @router.get("/api/connectors/types")
    def get_connector_types() -> dict:
        return {"items": list_connector_types()}

    @router.get("/api/connections")
    def list_connections() -> dict:
        return {"items": [public_connection_view(item) for item in ctx.repo.list_connections()]}

    @router.post("/api/connections")
    def create_connection(payload: ConnectionUpsertRequest) -> dict:
        connector_type = payload.connectorType.strip().lower()
        if connector_type not in {item["type"] for item in list_connector_types()}:
            raise HTTPException(status_code=400, detail="unsupported connector type")
        name = payload.name.strip()
        if not name:
            raise HTTPException(status_code=400, detail="name is required")
        credentials = payload.credentials or {}
        connection = {
            "name": name,
            "connectorType": connector_type,
            "config": payload.config or {},
            "credentialsEncrypted": encrypt_credentials(credentials, ctx.credentials_secret) if credentials else "",
        }
        saved = ctx.repo.save_connection(connection)
        audit(ctx.repo, "connection.create", saved["id"], {"connectorType": connector_type, "name": name})
        return public_connection_view(saved)

    @router.patch("/api/connections/{connection_id}")
    def update_connection(connection_id: str, payload: ConnectionUpsertRequest) -> dict:
        existing = ctx.repo.get_connection(connection_id)
        if not existing:
            raise HTTPException(status_code=404, detail="connection not found")
        connector_type = (payload.connectorType or existing.get("connectorType") or "").strip().lower()
        if connector_type not in {item["type"] for item in list_connector_types()}:
            raise HTTPException(status_code=400, detail="unsupported connector type")
        name = payload.name.strip() or existing.get("name") or connection_id
        next_connection = {
            **existing,
            "name": name,
            "connectorType": connector_type,
            "config": payload.config if payload.config is not None else existing.get("config") or {},
        }
        if payload.credentials is not None:
            next_connection["credentialsEncrypted"] = (
                encrypt_credentials(payload.credentials, ctx.credentials_secret) if payload.credentials else ""
            )
        saved = ctx.repo.save_connection(next_connection)
        audit(ctx.repo, "connection.update", connection_id, {"connectorType": connector_type, "name": name})
        return public_connection_view(saved)

    @router.delete("/api/connections/{connection_id}")
    def delete_connection(connection_id: str, request: Request) -> dict:
        ctx.require_admin_access(request)
        deleted = ctx.repo.delete_connection(connection_id)
        if not deleted:
            raise HTTPException(status_code=404, detail="connection not found")
        audit(ctx.repo, "connection.delete", connection_id)
        return {"deleted": True, "id": connection_id}

    @router.post("/api/connections/{connection_id}/test")
    def test_saved_connection(connection_id: str, payload: ConnectionTestRequest | None = None) -> dict:
        connection = ctx.repo.get_connection(connection_id)
        if not connection:
            raise HTTPException(status_code=404, detail="connection not found")
        body = payload or ConnectionTestRequest()
        connector_type = (body.connectorType or connection.get("connectorType") or "").strip().lower()
        config = body.config if body.config is not None else (connection.get("config") or {})
        if body.credentials is not None:
            credentials = body.credentials
        elif connection.get("credentialsEncrypted"):
            try:
                credentials = decrypt_credentials(connection["credentialsEncrypted"], ctx.credentials_secret)
            except ConnectorError as error:
                raise HTTPException(status_code=400, detail=error.message) from error
        else:
            credentials = {}
        try:
            result = test_connection(connector_type, config, credentials)
        except ConnectorError as error:
            connection["lastTestedAt"] = datetime.now(timezone.utc).isoformat()
            connection["lastTestStatus"] = "failed"
            ctx.repo.save_connection(connection)
            raise HTTPException(status_code=400, detail=error.message) from error
        connection["lastTestedAt"] = datetime.now(timezone.utc).isoformat()
        connection["lastTestStatus"] = "ok"
        ctx.repo.save_connection(connection)
        audit(ctx.repo, "connection.test", connection_id, {"ok": True, "rowCount": result.get("rowCount")})
        return result

    @router.post("/api/connections/test")
    def test_ephemeral_connection(payload: ConnectionTestRequest) -> dict:
        connector_type = (payload.connectorType or "").strip().lower()
        if not connector_type:
            raise HTTPException(status_code=400, detail="connectorType is required")
        try:
            return test_connection(connector_type, payload.config or {}, payload.credentials or {})
        except ConnectorError as error:
            raise HTTPException(status_code=400, detail=error.message) from error

    @router.post("/api/connections/{connection_id}/import")
    def import_connection_dataset(connection_id: str, payload: ConnectionImportRequest | None = None) -> dict:
        connection = ctx.repo.get_connection(connection_id)
        if not connection:
            raise HTTPException(status_code=404, detail="connection not found")
        body = payload or ConnectionImportRequest()
        credentials = {}
        if connection.get("credentialsEncrypted"):
            try:
                credentials = decrypt_credentials(connection["credentialsEncrypted"], ctx.credentials_secret)
            except ConnectorError as error:
                raise HTTPException(status_code=400, detail=error.message) from error
        try:
            scraped_payload = fetch_connector_payload(
                connection.get("connectorType") or "",
                connection.get("config") or {},
                credentials,
            )
            data_source = build_connector_data_source(
                name=body.datasetName or connection.get("name") or connection_id,
                connector_type=str(connection.get("connectorType") or "connector"),
                payload=scraped_payload,
                connection_id=connection_id,
            )
        except ConnectorError as error:
            raise HTTPException(status_code=400, detail=error.message) from error
        data_source["uploadedAt"] = datetime.now(timezone.utc).isoformat()
        dataset = ctx.repo.add_uploaded_dataset(data_source)
        connection["lastImportAt"] = datetime.now(timezone.utc).isoformat()
        ctx.repo.save_connection(connection)
        audit(
            ctx.repo,
            "connection.import",
            connection_id,
            {
                "datasetId": dataset.get("id"),
                "kind": dataset.get("kind"),
                "rowCount": dataset.get("rowCount"),
                "columnCount": dataset.get("columnCount"),
            },
        )
        return dataset

    def _run_model_complete(prompt: str) -> str:
        client = getattr(ctx, "client", None)
        if client is None or not hasattr(client, "complete"):
            raise RuntimeError("model client unavailable")
        result = client.complete(prompt)
        if inspect.isawaitable(result):
            try:
                asyncio.get_running_loop()
            except RuntimeError:
                return str(asyncio.run(result) or "")
            raise RuntimeError("async model client cannot be awaited in sync route")
        return str(result or "")

    @router.post("/api/smart-scrape")
    def smart_scrape_endpoint(payload: SmartScrapeRequest) -> dict:
        """Natural-language guided web scrape.

        - commit=false: preview only (no dataset import)
        - commit=true : generate file + import dataset (legacy/default)
        """
        url = (payload.url or "").strip()
        instruction = (payload.instruction or "").strip()
        if not url:
            raise HTTPException(status_code=400, detail="url is required")
        if not instruction:
            raise HTTPException(status_code=400, detail="instruction is required")

        credentials: dict = {}
        if payload.bearerToken:
            credentials["bearerToken"] = payload.bearerToken

        model_complete = None
        if getattr(ctx, "api_key", None) and getattr(ctx, "client", None) is not None:
            model_complete = _run_model_complete

        try:
            result = smart_scrape_to_dataset(
                url=url,
                instruction=instruction,
                dataset_name=payload.datasetName,
                credentials=credentials,
                model_complete=model_complete,
                preferred_format=payload.outputFormat,
                allow_private_hosts=bool(payload.allowPrivateHosts),
                render_js=bool(getattr(payload, "renderJs", False)),
            )
        except ConnectorError as error:
            raise HTTPException(
                status_code=400,
                detail={"code": getattr(error, "code", "connector_error"), "message": error.message},
            ) from error

        data_source = result["dataset"]
        file_info = result["file"]
        extraction = result.get("extraction") or {}
        preview_rows = extraction.get("preview") or []
        if not isinstance(preview_rows, list):
            preview_rows = []

        # Preview mode: do not persist dataset yet.
        if not bool(payload.commit):
            audit(
                ctx.repo,
                "smart_scrape.preview",
                "preview",
                {
                    "url": url,
                    "instruction": instruction[:200],
                    "kind": extraction.get("kind"),
                    "format": file_info.get("format"),
                    "strategy": (extraction.get("meta") or {}).get("strategy"),
                },
            )
            return {
                "committed": False,
                "message": "已生成采集预览，请确认后再导入。",
                "file": {
                    "fileName": file_info.get("fileName"),
                    "mimeType": file_info.get("mimeType"),
                    "format": file_info.get("format"),
                    "byteSize": file_info.get("byteSize"),
                    "contentBase64": file_info.get("contentBase64"),
                    "contentTextPreview": str(file_info.get("contentText") or "")[:2000],
                },
                "extraction": {
                    "kind": extraction.get("kind"),
                    "meta": extraction.get("meta") or {},
                    "preview": preview_rows[:8],
                    "previewCount": len(preview_rows),
                    # Pass structured payload back for commit without re-scrape.
                    "rows": data_source.get("analysisRows") if data_source.get("kind") == "table" else None,
                    "text": data_source.get("fullText") or data_source.get("textPreview"),
                    "paragraphs": [
                        str((row or {}).get("text") or "")
                        for row in (data_source.get("analysisRows") or [])
                        if isinstance(row, dict)
                    ]
                    if data_source.get("kind") == "document"
                    else None,
                },
                "suggestedDatasetName": payload.datasetName
                or data_source.get("name")
                or file_info.get("fileName"),
            }

        data_source["uploadedAt"] = datetime.now(timezone.utc).isoformat()

        stored = None
        try:
            raw = base64.b64decode(file_info["contentBase64"].encode("ascii"))
            tenant_id, _owner_id = get_workspace_scope()
            stored = ctx.store.save("smart-scrape", file_info["fileName"], raw, tenant_id=tenant_id)
            data_source["storageKey"] = stored.get("storageKey")
            data_source["originalFileSize"] = stored.get("byteSize") or file_info.get("byteSize")
            data_source["originalMimeType"] = file_info.get("mimeType")
        except Exception:
            data_source["originalFileSize"] = file_info.get("byteSize")
            data_source["originalMimeType"] = file_info.get("mimeType")

        dataset = ctx.repo.add_uploaded_dataset(data_source)

        active_task = None
        if payload.startAnalysis:
            active_task = ctx.repo.update_active_task(
                {
                    "dataSource": dataset,
                    "status": "等待需求解析",
                    "request": instruction,
                }
            )

        audit(
            ctx.repo,
            "smart_scrape.import",
            dataset.get("id") or "dataset",
            {
                "url": url,
                "instruction": instruction[:200],
                "kind": dataset.get("kind"),
                "format": file_info.get("format"),
                "strategy": (extraction.get("meta") or {}).get("strategy"),
                "startAnalysis": bool(payload.startAnalysis),
            },
        )
        return {
            "committed": True,
            "dataset": dataset,
            "file": {
                "fileName": file_info.get("fileName"),
                "mimeType": file_info.get("mimeType"),
                "format": file_info.get("format"),
                "byteSize": file_info.get("byteSize"),
                "downloadUrl": (
                    f"/api/datasets/{dataset.get('id')}/download"
                    if dataset.get("id") and dataset.get("storageKey")
                    else None
                ),
                "storageKey": (stored or {}).get("storageKey"),
            },
            "extraction": extraction,
            "activeTask": active_task,
            "message": "已按指令抓取内容并生成文件，已导入为可分析数据源。",
        }

    @router.post("/api/smart-scrape/commit")
    def smart_scrape_commit_endpoint(payload: SmartScrapeCommitRequest) -> dict:
        """Commit a previewed scrape payload into dataset storage without re-fetching the page."""
        url = (payload.url or "").strip()
        instruction = (payload.instruction or "").strip()
        if not url or not instruction:
            raise HTTPException(status_code=400, detail="url and instruction are required")
        if not payload.contentBase64 or not payload.fileName:
            raise HTTPException(status_code=400, detail="file content is required")

        kind = (payload.kind or "").strip().lower()
        rows = payload.rows if isinstance(payload.rows, list) else None
        text = (payload.text or "").strip()
        paragraphs = payload.paragraphs if isinstance(payload.paragraphs, list) else None
        selected_fields = [str(item).strip() for item in (payload.selectedFields or []) if str(item).strip()]
        selected_row_indexes = [
            int(item)
            for item in (payload.selectedRowIndexes or [])
            if str(item).strip().lstrip("-").isdigit() and int(item) >= 0
        ]

        # Project selected rows first, then fields.
        if kind == "table" and rows and selected_row_indexes:
            projected_rows: list[dict] = []
            for index in selected_row_indexes:
                if index < len(rows) and isinstance(rows[index], dict):
                    projected_rows.append(rows[index])
            rows = projected_rows

        if kind == "table" and rows and selected_fields:
            projected: list[dict] = []
            for row in rows:
                if not isinstance(row, dict):
                    continue
                projected.append({field: row.get(field, "") for field in selected_fields})
            rows = projected

        content_base64 = payload.contentBase64
        file_name = payload.fileName
        mime_type = payload.mimeType
        file_format = payload.format

        if kind == "table" and rows:
            # Rebuild file content from projected rows so download matches imported dataset.
            try:
                from backend_py.smart_scrape import materialize_scrape_file

                rebuilt = materialize_scrape_file(
                    payload={"kind": "table", "rows": rows, "meta": payload.extractionMeta or {}},
                    instruction=instruction,
                    source_url=url,
                    preferred_format=file_format or "csv",
                )
                content_base64 = rebuilt["contentBase64"]
                file_name = rebuilt["fileName"]
                mime_type = rebuilt["mimeType"]
                file_format = rebuilt["format"]
            except Exception:
                pass
            extracted = {
                "kind": "table",
                "rows": rows,
                "meta": {
                    **(payload.extractionMeta or {}),
                    "sourceUrl": url,
                    "instruction": instruction,
                    "strategy": (payload.extractionMeta or {}).get("strategy") or "preview_commit",
                    "selectedFields": selected_fields,
                    "selectedRowIndexes": selected_row_indexes,
                },
            }
        else:
            body_text = text
            if not body_text and paragraphs:
                body_text = "\n\n".join(str(item) for item in paragraphs if str(item).strip())
            if not body_text:
                # Fallback: decode generated file text.
                try:
                    body_text = base64.b64decode(content_base64.encode("ascii")).decode("utf-8", errors="replace")
                except Exception as error:
                    raise HTTPException(status_code=400, detail="invalid file content") from error
            extracted = {
                "kind": "document",
                "text": body_text[:200_000],
                "paragraphs": paragraphs or [item for item in body_text.split("\n\n") if item.strip()][:200],
                "meta": {
                    **(payload.extractionMeta or {}),
                    "sourceUrl": url,
                    "instruction": instruction,
                    "strategy": (payload.extractionMeta or {}).get("strategy") or "preview_commit",
                },
            }

        try:
            data_source = build_connector_data_source(
                name=(payload.datasetName or file_name).strip() or file_name,
                connector_type="smart_web",
                payload=extracted,
                connection_id=None,
            )
        except ConnectorError as error:
            raise HTTPException(status_code=400, detail=error.message) from error

        data_source["name"] = (payload.datasetName or file_name).strip() or file_name
        data_source["fileName"] = file_name
        data_source["formatLabel"] = f"SmartScrape:{(file_format or 'bin').upper()}"
        data_source["sourceType"] = "smart_scrape"
        data_source["connectorType"] = "smart_web"
        data_source["sourceUrl"] = url
        data_source["scrapeInstruction"] = instruction
        data_source["scrapeMeta"] = extracted.get("meta") or {}
        data_source["uploadedAt"] = datetime.now(timezone.utc).isoformat()
        data_source["originalMimeType"] = mime_type
        data_source["generatedFile"] = {
            "fileName": file_name,
            "mimeType": mime_type,
            "format": file_format,
            "byteSize": 0,
            "contentBase64": content_base64,
        }

        stored = None
        try:
            raw = base64.b64decode(content_base64.encode("ascii"))
            tenant_id, _owner_id = get_workspace_scope()
            stored = ctx.store.save("smart-scrape", file_name, raw, tenant_id=tenant_id)
            data_source["storageKey"] = stored.get("storageKey")
            data_source["originalFileSize"] = stored.get("byteSize") or len(raw)
            data_source["generatedFile"]["byteSize"] = data_source["originalFileSize"]
        except Exception:
            try:
                raw = base64.b64decode(content_base64.encode("ascii"))
                data_source["originalFileSize"] = len(raw)
            except Exception:
                data_source["originalFileSize"] = 0

        dataset = ctx.repo.add_uploaded_dataset(data_source)
        active_task = None
        if payload.startAnalysis:
            active_task = ctx.repo.update_active_task(
                {
                    "dataSource": dataset,
                    "status": "等待需求解析",
                    "request": instruction,
                }
            )

        audit(
            ctx.repo,
            "smart_scrape.commit",
            dataset.get("id") or "dataset",
            {
                "url": url,
                "instruction": instruction[:200],
                "kind": dataset.get("kind"),
                "format": file_format,
                "selectedFields": selected_fields,
                "selectedRowCount": len(selected_row_indexes) if selected_row_indexes else None,
                "startAnalysis": bool(payload.startAnalysis),
            },
        )
        return {
            "committed": True,
            "dataset": dataset,
            "file": {
                "fileName": file_name,
                "mimeType": mime_type,
                "format": file_format,
                "byteSize": data_source.get("originalFileSize"),
                "downloadUrl": (
                    f"/api/datasets/{dataset.get('id')}/download"
                    if dataset.get("id") and dataset.get("storageKey")
                    else None
                ),
                "storageKey": (stored or {}).get("storageKey"),
            },
            "activeTask": active_task,
            "message": "已确认导入采集结果。",
        }

    @router.delete("/api/datasets/{dataset_id}")
    def delete_dataset(dataset_id: str, request: Request) -> dict:
        ctx.require_admin_access(request)
        if ctx.repo.active_task_uses_dataset(dataset_id):
            raise HTTPException(status_code=409, detail="active task is using this dataset")
        dataset = ctx.repo.delete_dataset(dataset_id)
        if not dataset:
            raise HTTPException(status_code=404, detail="dataset not found")
        storage_key = dataset.get("storageKey")
        if storage_key:
            tenant_id, _owner_id = get_workspace_scope()
            try:
                ctx.store.delete(storage_key, tenant_id=tenant_id)
            except Exception:
                pass
        audit(ctx.repo, "dataset.delete", dataset_id)
        return {"deleted": True, "id": dataset_id}

    @router.get("/api/reports")
    def list_reports() -> list[dict]:
        return ctx.repo.list_reports()

    @router.get("/api/search")
    def search(query: str = "") -> list[dict]:
        return ctx.repo.search(query)

    @router.get("/api/notifications")
    def list_notifications() -> list[dict]:
        return ctx.repo.list_notifications()

    def persist_uploaded_data_source(file_name: str, mime_type: str | None, content: bytes) -> dict:
        try:
            checked = validate_upload(
                file_name=file_name,
                content_type=mime_type,
                size_bytes=len(content),
                max_bytes=ctx.max_request_body_bytes,
                content=content,
            )
        except UploadSecurityError as error:
            raise HTTPException(status_code=400, detail=error.message) from error
        encoded = base64.b64encode(content).decode("ascii")
        try:
            parsed = ingest_data_source(checked["fileName"], checked["contentType"], encoded)
        except ValueError as error:
            raise HTTPException(status_code=400, detail=str(error)) from error
        tenant_id, _owner_id = get_workspace_scope()
        stored = ctx.store.save("uploads", checked["fileName"], content, tenant_id=tenant_id)
        parsed["storageKey"] = stored["storageKey"]
        parsed["originalMimeType"] = checked["contentType"]
        parsed["originalFileSize"] = len(content)
        parsed["uploadedAt"] = datetime.now(timezone.utc).isoformat()
        parsed["storageBackend"] = stored.get("backend") or ctx.artifact_backend
        dataset = ctx.repo.add_uploaded_dataset(parsed)
        audit(
            ctx.repo,
            "data_source.ingest",
            dataset["id"],
            {
                "kind": parsed["kind"],
                "formatLabel": parsed["formatLabel"],
                "rowCount": parsed["rowCount"],
                "storageBackend": parsed.get("storageBackend"),
            },
        )
        return dataset

    @router.post("/api/data-source/upload")
    async def upload_data_source(file: UploadFile = File(...)) -> dict:
        content = await file.read()
        if len(content) > ctx.max_request_body_bytes:
            raise HTTPException(status_code=413, detail="uploaded file exceeds request size limit")
        return persist_uploaded_data_source(file.filename or "upload", file.content_type, content)

    @router.post("/api/data-source/ingest")
    def ingest_uploaded_data_source(payload: DataSourceIngestRequest) -> dict:
        try:
            content = base64.b64decode(payload.contentBase64, validate=True)
        except ValueError as error:
            raise HTTPException(status_code=400, detail="invalid file payload") from error
        if len(content) > ctx.max_request_body_bytes:
            raise HTTPException(status_code=413, detail="request body too large")
        try:
            return persist_uploaded_data_source(payload.fileName, payload.mimeType, content)
        except HTTPException as error:
            raise error

    @router.get("/api/datasets/{dataset_id}/download")
    def download_dataset(dataset_id: str):
        dataset = ctx.repo.get_dataset(dataset_id)
        if not dataset or not dataset.get("storageKey"):
            raise HTTPException(status_code=404, detail="dataset file not found")
        tenant_id, _owner_id = get_workspace_scope()
        try:
            path = ctx.store.resolve(dataset["storageKey"], tenant_id=tenant_id)
        except PermissionError as error:
            raise HTTPException(status_code=403, detail="dataset access denied") from error
        except (OSError, ValueError) as error:
            raise HTTPException(status_code=404, detail="dataset file not found") from error
        return FileResponse(path, media_type=dataset.get("originalMimeType") or "application/octet-stream", filename=dataset["name"])

    @router.get("/api/tasks/active")
    def get_canonical_active_task() -> dict:
        tenant_id, owner_id = get_workspace_scope()
        job_view = ctx.workflow_runner.job_public_view(ctx.workflow_runner.get_job(tenant_id, owner_id))
        return attach_execution_meta(ctx.repo.get_active_task(), job_view, mode="async" if ctx.workflow_async else "sync")

    @router.post("/api/tasks/active/dataset/{dataset_id}")
    def attach_dataset_to_active_task(dataset_id: str) -> dict:
        dataset = ctx.repo.get_dataset(dataset_id)
        if not dataset:
            raise HTTPException(status_code=404, detail="dataset not found")
        task = ctx.repo.update_active_task({"dataSource": dataset, "status": "等待需求解析"})
        audit(ctx.repo, "task.dataset.attach", task["id"], {"datasetId": dataset_id})
        return task

    @router.delete("/api/tasks/active/report")
    def delete_active_report(request: Request) -> dict:
        ctx.require_admin_access(request)
        tenant_id, _owner_id = get_workspace_scope()
        for artifact in ctx.repo.delete_artifacts_for_task("workspace-active-session"):
            storage_key = artifact.get("storageKey")
            if storage_key:
                try:
                    ctx.store.delete(storage_key, tenant_id=tenant_id)
                except Exception:
                    pass
        task = ctx.repo.reset_active_task()
        audit(ctx.repo, "active_task.report_delete", task["id"])
        return {"deleted": True, "id": task["id"]}

    @router.post("/api/tasks/active/correction")
    def apply_task_correction(payload: TaskCorrectionRequest) -> dict:
        task = ctx.repo.get_active_task()
        snapshot = task["snapshot"]
        nodes = [{**node, "status": "success" if node["id"] in {"source", "clean"} else "pending"} for node in snapshot["nodes"]]
        snapshot["nodes"] = nodes
        snapshot["phase"] = "idle"
        snapshot["logMessage"] = "已记录人工修正，指标及后续节点等待重新执行"
        snapshot["context"] = {**snapshot["context"], "latestLog": snapshot["logMessage"]}
        task = ctx.repo.update_active_task({"summary": payload.summary, "snapshot": snapshot, "status": "等待重新执行", "artifacts": [item for item in task.get("artifacts", []) if item.get("nodeId") in {"source", "clean"}]})
        audit(ctx.repo, "task.correction", task["id"])
        return task

    @router.post("/api/tasks/active/exports")
    def export_active_task(payload: ExportRequest, request: Request) -> dict:
        task = ctx.repo.get_active_task()
        audit_artifact = next((item for item in reversed(task.get("artifacts", [])) if item.get("type") == "audit"), None)
        if not audit_artifact:
            raise HTTPException(status_code=409, detail="complete workflow audit before exporting")
        report = payload.report or next(
            (
                item.get("content", {}).get("markdown")
                for item in reversed(task.get("artifacts", []))
                if item.get("type") == "report_markdown"
            ),
            task.get("summary", ""),
        )
        if ctx.export_mask_sensitive:
            report = mask_text(str(report or ""))
        auth_user = getattr(request.state, "auth_user", None)
        tenant_id, owner_id = get_workspace_scope()
        watermark_lines = build_watermark_lines(
            username=(auth_user.username if auth_user else owner_id),
            tenant_id=(auth_user.tenant_id if auth_user else tenant_id),
        )
        if ctx.export_watermark_enabled:
            report = apply_text_watermark(str(report or ""), watermark_lines)
        export_format = payload.format.lower()
        safe_name = "".join(
            character if character not in '\\/:*?"<>|' else "_"
            for character in task.get("name", "analysis-report")
        )
        if export_format in {"markdown", "md"}:
            file_name, mime_type, content = (
                f"{safe_name}.md",
                "text/markdown; charset=utf-8",
                str(report).encode("utf-8"),
            )
        elif export_format == "html":
            html = (
                f"<!doctype html><meta charset='utf-8'><title>{escape(task.get('name', '分析报告'))}</title>"
                f"<pre>{escape(str(report))}</pre>"
            )
            if ctx.export_watermark_enabled:
                html = apply_html_watermark(html, watermark_lines)
            file_name, mime_type, content = (
                f"{safe_name}.html",
                "text/html; charset=utf-8",
                html.encode("utf-8"),
            )
        elif export_format in {"word", "docx"}:
            from docx import Document

            document = Document()
            document.add_heading(task.get("name", "分析报告"), 0)
            if ctx.export_watermark_enabled:
                for line in watermark_lines:
                    document.add_paragraph(line)
            for line in str(report).splitlines():
                document.add_paragraph(line)
            temp = ctx.store.root / "exports" / f"{uuid4()}.docx"
            temp.parent.mkdir(parents=True, exist_ok=True)
            document.save(temp)
            content = temp.read_bytes()
            temp.unlink(missing_ok=True)
            file_name, mime_type = (
                f"{safe_name}.docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        elif export_format == "pdf":
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.pdfgen import canvas

            pdfctx.metrics.registerFont(UnicodeCIDFont("STSong-Light"))
            temp = ctx.store.root / "exports" / f"{uuid4()}.pdf"
            temp.parent.mkdir(parents=True, exist_ok=True)
            pdf = canvas.Canvas(str(temp), pagesize=A4)
            width, height = A4
            if ctx.export_watermark_enabled:
                pdf.saveState()
                pdf.setFont("STSong-Light", 18)
                pdf.setFillGray(0.88)
                pdf.translate(width / 2, height / 2)
                pdf.rotate(30)
                pdf.drawCentredString(0, 0, watermark_lines[0][:40])
                pdf.restoreState()
            pdf.setFont("STSong-Light", 11)
            y = 800
            for line in str(report).splitlines() or ["分析报告"]:
                pdf.drawString(48, y, line[:90])
                y -= 18
                if y < 48:
                    pdf.showPage()
                    if ctx.export_watermark_enabled:
                        pdf.saveState()
                        pdf.setFont("STSong-Light", 18)
                        pdf.setFillGray(0.88)
                        pdf.translate(width / 2, height / 2)
                        pdf.rotate(30)
                        pdf.drawCentredString(0, 0, watermark_lines[0][:40])
                        pdf.restoreState()
                    pdf.setFont("STSong-Light", 11)
                    y = 800
            pdf.save()
            content = temp.read_bytes()
            temp.unlink(missing_ok=True)
            file_name, mime_type = f"{safe_name}.pdf", "application/pdf"
        else:
            raise HTTPException(status_code=400, detail="unsupported export format")
        tenant_id, _owner_id = get_workspace_scope()
        stored = ctx.store.save("exports", file_name, content, tenant_id=tenant_id)
        artifact = {
            **stored,
            "taskId": task["id"],
            "type": "export",
            "format": export_format,
            "mimeType": mime_type,
            "createdAt": datetime.now(timezone.utc).isoformat(),
            "watermark": ctx.export_watermark_enabled,
            "masked": ctx.export_mask_sensitive,
        }
        ctx.repo.add_artifact(artifact)
        exports = [
            {
                "format": export_format.upper(),
                "status": f"{export_format.upper()} 导出完成",
                "artifactId": artifact["id"],
            },
            *task.get("exports", []),
        ]
        ctx.repo.update_active_task({"exports": exports})
        audit(
            ctx.repo,
            "task.export",
            task["id"],
            {
                "artifactId": artifact["id"],
                "format": export_format,
                "watermark": ctx.export_watermark_enabled,
                "masked": ctx.export_mask_sensitive,
            },
        )
        return {**artifact, "downloadUrl": f"/api/artifacts/{artifact['id']}/download"}

    @router.get("/api/artifacts/{artifact_id}/download")
    def download_artifact(artifact_id: str):
        artifact = ctx.repo.get_artifact(artifact_id)
        if not artifact:
            raise HTTPException(status_code=404, detail="artifact not found")
        tenant_id, _owner_id = get_workspace_scope()
        try:
            path = ctx.store.resolve(artifact["storageKey"], tenant_id=tenant_id)
        except PermissionError as error:
            raise HTTPException(status_code=403, detail="artifact access denied") from error
        except (OSError, ValueError) as error:
            raise HTTPException(status_code=404, detail="artifact file not found") from error
        return FileResponse(path, media_type=artifact.get("mimeType") or "application/octet-stream", filename=artifact.get("fileName") or "artifact")

    @router.post("/api/templates/save-active-task")
    def save_active_task_as_template(payload: TemplateDraftRequest, request: Request) -> dict:
        ctx.require_admin_access(request)
        template = ctx.repo.save_active_task_as_template(payload.model_dump(exclude_none=True))
        audit(ctx.repo, "template.save_active_task", template["id"])
        return template

    @router.get("/api/audit/events")
    def list_audit_events(request: Request, limit: int = 50, offset: int = 0) -> dict:
        ctx.require_admin_access(request)
        if not ctx.enable_audit_events:
            raise HTTPException(status_code=404, detail="Not found")
        safe_limit = max(1, min(limit, 200))
        safe_offset = max(0, offset)
        total = ctx.repo.count_audit_events()
        items = ctx.repo.list_audit_events(limit=safe_limit, offset=safe_offset)
        return {
            "items": items,
            "pagination": {
                "limit": safe_limit,
                "offset": safe_offset,
                "total": total,
                "hasMore": safe_offset + len(items) < total,
            },
        }

    @router.post("/api/intent/parse")
    async def parse_intent(payload: IntentRequest):
        request_text = payload.request.strip()
        if not request_text:
            raise HTTPException(status_code=400, detail="request is required")

        prompt = "\n".join(
            [
                "Convert the analytics request below into JSON only.",
                "Return only JSON.",
                "Use keys: request, goal, metrics, dimensions, charts.",
                f"Request: {request_text}{data_source_context(payload.dataSource)}",
            ]
        )

        try:
            content = await ctx.client.complete(prompt)
            model_payload = json.loads(content)
            if not isinstance(model_payload, dict):
                model_payload = {}
            return build_agentic_intent(request_text, payload.dataSource, model_payload)
        except Exception:
            return build_agentic_intent(request_text, payload.dataSource)


    app.include_router(router)
